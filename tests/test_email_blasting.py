"""Checks for the supplied workbook layout using invented data only."""

import io
from datetime import datetime, timezone
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from EmailBlasting_Template import (
    OPTIONAL_COLUMNS, REQUIRED_COLUMNS, TEMPLATE_COLUMNS, ValidationError,
    export_workbook, output_filename, populate_template, prepare_group_downloads,
    read_source, split_amnesty, suggest_mapping,
)


APP_PATH = Path(__file__).resolve().parents[1] / "EmailBlasting_Template.py"
SOURCE_CSV = (
    b"Email,Primary Key,Index One UID,Card Number,CH Name,Phone Number,Outstanding Balance,Agent,Code,DO Class,Level,Product,Endo Date\n"
    b"alex@example.com,0001,000099,0000000000000001,Alex Example,09123456789,1234.50,Agent One,NORTH,A1,L1,Example Product,2026-10-13\n"
    b"alex@example.com,0002,000100,0000000000000002,Jamie Example,09987654321,0,Agent Two,SOUTH,A2,L2,Other Product,2026-10-14\n"
)


class TemplateTests(unittest.TestCase):
    def test_revised_template_maps_required_fields_with_optional_columns_absent(self):
        source = read_source(SOURCE_CSV, "source.csv")
        mapping = suggest_mapping(source.columns)
        self.assertEqual(set(mapping), set(TEMPLATE_COLUMNS))
        self.assertTrue(all(mapping[field] for field in REQUIRED_COLUMNS))
        self.assertTrue(all(mapping[field] is None for field in ("ChristmasAmnesty", "NewEndo")))
        output = populate_template(source, mapping)
        self.assertEqual(output.columns.tolist(), TEMPLATE_COLUMNS)
        self.assertEqual(output["ChristmasAmnesty"].tolist(), ["", ""])
        self.assertEqual(output["NewEndo"].tolist(), ["", ""])
        self.assertEqual(output["primary_key"].tolist(), ["0001", "0002"])
        self.assertEqual(output["phone"].tolist(), ["09123456789", "09987654321"])
        self.assertEqual(output["cardnumber"].tolist(), ["0000000000000001", "0000000000000002"])
        self.assertEqual(output["email"].tolist(), ["alex@example.com", "alex@example.com"])

    def test_optional_columns_copy_values_and_allow_blank_cells(self):
        source = read_source(SOURCE_CSV, "source.csv")
        source["Christmas Amnesty"] = ["Eligible", ""]
        source["New Endo"] = ["", "2026-10-15"]
        output = populate_template(source, suggest_mapping(source.columns))
        self.assertEqual(output["ChristmasAmnesty"].tolist(), ["Eligible", ""])
        self.assertEqual(output["NewEndo"].tolist(), ["", "2026-10-15"])

    def test_optional_fixed_values_include_zero(self):
        source = read_source(SOURCE_CSV, "source.csv")
        output = populate_template(source, suggest_mapping(source.columns), {"ChristmasAmnesty": "Yes", "NewEndo": 0})
        self.assertEqual(output["ChristmasAmnesty"].tolist(), ["Yes", "Yes"])
        self.assertEqual(output["NewEndo"].tolist(), ["0", "0"])

    def test_one_optional_column_can_be_present_without_the_other(self):
        source = read_source(SOURCE_CSV, "source.csv")
        source["NewEndo"] = ["Renewed", ""]
        output = populate_template(source, suggest_mapping(source.columns))
        self.assertEqual(output["NewEndo"].tolist(), ["Renewed", ""])
        self.assertEqual(output["ChristmasAmnesty"].tolist(), ["", ""])

    def test_missing_data_field_requires_mapping_or_fixed_value(self):
        source = read_source(SOURCE_CSV, "source.csv").drop(columns=["Product"])
        mapping = suggest_mapping(source.columns)
        with self.assertRaisesRegex(ValidationError, "Product"):
            populate_template(source, mapping)
        output = populate_template(source, mapping, {"Product": "Example Product"})
        self.assertEqual(output["Product"].tolist(), ["Example Product", "Example Product"])

    def test_ambiguous_headers_are_not_silently_selected(self):
        mapping = suggest_mapping(["CH Name", "chname", "Email Address"])
        self.assertIsNone(mapping["chname"])
        self.assertEqual(mapping["email"], "Email Address")

    def test_excel_header_row_sheet_and_numeric_zero_padding(self):
        workbook = Workbook()
        workbook.active.title = "Notes"
        sheet = workbook.create_sheet("Contacts")
        sheet.append(["Synthetic title"])
        sheet.append([column for column in TEMPLATE_COLUMNS if column not in {"ChristmasAmnesty", "NewEndo"}])
        sheet.append(["alex@example.com", 12, "000099", "0000000000000001", "Alex Example", 9123456789, 1234.5, "Agent One", "NA", "A1", "L1", "Example Product", "2026-10-13"])
        sheet["B3"].number_format = "000000"
        sheet["F3"].number_format = "00000000000"
        buffer = io.BytesIO()
        workbook.save(buffer)
        source = read_source(buffer.getvalue(), "source.xlsx", "Contacts", header_row=2)
        self.assertEqual(source.iloc[0]["primary_key"], "000012")
        self.assertEqual(source.iloc[0]["phone"], "09123456789")
        self.assertEqual(source.iloc[0]["CODE"], "NA")

    def test_plain_output_preserves_values_and_treats_formulas_as_text(self):
        source = read_source(SOURCE_CSV, "source.csv")
        source.loc[0, "CH Name"] = "=2+2"
        source.loc[1, "Agent"] = "+123"
        output = populate_template(source, suggest_mapping(source.columns))
        workbook = load_workbook(io.BytesIO(export_workbook(output)), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Sheet1"])
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[1]], TEMPLATE_COLUMNS)
        self.assertEqual(sheet.max_column, 15)
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["B2"].value, "0001")
        self.assertEqual(sheet["F2"].value, "09123456789")
        self.assertEqual(sheet["E2"].value, "=2+2")
        self.assertEqual(sheet["E2"].data_type, "s")
        self.assertEqual(sheet["H3"].data_type, "s")
        self.assertIsNone(sheet["N2"].value)
        self.assertIsNone(sheet["O2"].value)
        self.assertEqual(sheet["M1"].value, "EndoDate")
        self.assertEqual(sheet["N1"].value, "ChristmasAmnesty")
        self.assertEqual(sheet["O1"].value, "NewEndo")
        self.assertFalse(any(char.isspace() for cell in sheet[1] for char in cell.value))
        self.assertFalse(sheet.column_dimensions)
        self.assertFalse(sheet.row_dimensions)
        self.assertFalse(sheet.tables)
        self.assertTrue(all(not cell.has_style for row in sheet for cell in row))

    def test_excel_without_sheet_dimensions_can_be_imported_and_is_row_limited(self):
        source = read_source(SOURCE_CSV, "source.csv")
        output = populate_template(source, suggest_mapping(source.columns))
        content = export_workbook(output)
        workbook = load_workbook(io.BytesIO(content), read_only=True)
        self.assertIsNone(workbook.active.max_row)
        workbook.close()
        restored = read_source(content, "export.xlsx")
        self.assertEqual(restored["phone"].tolist(), output["phone"].tolist())
        self.assertEqual(restored["primary_key"].tolist(), output["primary_key"].tolist())
        with patch("EmailBlasting_Template.MAX_SOURCE_ROWS", 1):
            with self.assertRaisesRegex(ValidationError, "1 data rows"):
                read_source(content, "export.xlsx")

    def test_bad_csv_and_duplicate_headers_are_rejected(self):
        for content in (b"email,email\na@example.com,b@example.com\n", b"email,name\na@example.com,Alex,Extra\n", b"email,name\n"):
            with self.assertRaises(ValidationError):
                read_source(content, "source.csv")

    def test_blank_rows_and_unused_columns_are_ignored(self):
        source = read_source(b"email,Name,\nalex@example.com,Alex,\n,,\n", "source.csv")
        self.assertEqual(source.to_dict("records"), [{"email": "alex@example.com", "Name": "Alex"}])
        with self.assertRaises(ValidationError):
            read_source(b"email,,\nalex@example.com,Missing header,\n", "source.csv")

    def test_filename_identifies_area_source_and_manila_export_time(self):
        exported_at = datetime(2026, 10, 13, 12, 30, 45, 123456, tzinfo=timezone.utc)
        self.assertEqual(
            output_filename("NORTH", "source.csv", exported_at),
            "EmailBlasting_(NORTH)_source_20261013_203045_123456.xlsx",
        )
        self.assertEqual(
            output_filename("North/West", r"C:\Reports\Source File.xlsx", exported_at),
            "EmailBlasting_(North_West)_Source_File_20261013_203045_123456.xlsx",
        )


class InterfaceTests(unittest.TestCase):
    def start_app(self):
        app = AppTest.from_file(str(APP_PATH), default_timeout=20).run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        return app

    def test_upload_generate_and_clear(self):
        app = self.start_app()
        app.file_uploader[0].set_value(("source.csv", SOURCE_CSV, "text/csv"))
        app.run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertEqual(app.dataframe[-1].value.columns.tolist(), TEMPLATE_COLUMNS)
        self.assertEqual(len(app.selectbox), 15)
        for field in ("ChristmasAmnesty", "NewEndo"):
            widget = next(widget for widget in app.selectbox if widget.label.startswith(field))
            self.assertEqual(widget.value, "Leave blank")
        self.assertIn("email_result", app.session_state)
        self.assertEqual(len(app.get("download_button")), 1)
        workbook = load_workbook(io.BytesIO(app.session_state["email_result"]))
        self.assertEqual([cell.value for cell in workbook.active[1]], TEMPLATE_COLUMNS)
        next(button for button in app.button if button.label == "Clear source file").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.file_uploader[0].value)
        self.assertNotIn("email_result", app.session_state)
        self.assertNotIn("email_source", app.session_state)
        self.assertNotIn("email_result_created_at", app.session_state)
        self.assertFalse(app.get("download_button"))

    def test_mapping_change_refreshes_the_download_values(self):
        app = self.start_app()
        app.file_uploader[0].set_value(("source.csv", SOURCE_CSV, "text/csv"))
        app.run()
        self.assertEqual(len(app.get("download_button")), 1)
        next(widget for widget in app.selectbox if widget.label.startswith("Product")).set_value("Use a fixed value").run()
        self.assertFalse(app.get("download_button"))
        next(widget for widget in app.text_input if widget.label == "Product value for all rows").set_value("Updated Product").run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertEqual(len(app.get("download_button")), 1)
        self.assertEqual(app.dataframe[-1].value["Product"].tolist(), ["Updated Product", "Updated Product"])
        workbook = load_workbook(io.BytesIO(app.session_state["email_result"]))
        self.assertEqual(workbook.active["L2"].value, "Updated Product")
        self.assertEqual(workbook.active["L3"].value, "Updated Product")

    def test_uploaded_data_is_session_specific(self):
        first = self.start_app()
        second = self.start_app()
        first.file_uploader[0].set_value(("source.csv", SOURCE_CSV, "text/csv"))
        first.run()
        self.assertNotIn("email_source", second.session_state)
        self.assertFalse(second.dataframe)

    def test_optional_mapping_can_be_copied_or_explicitly_left_blank(self):
        source = read_source(SOURCE_CSV, "source.csv")
        source["Christmas Amnesty"] = ["Eligible", ""]
        source["New Endo"] = ["", "Renewed"]
        app = self.start_app()
        app.file_uploader[0].set_value(("source.csv", source.to_csv(index=False).encode(), "text/csv"))
        app.run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertEqual(app.dataframe[-1].value["ChristmasAmnesty"].tolist(), ["Eligible", ""])
        next(widget for widget in app.selectbox if widget.label.startswith("ChristmasAmnesty")).set_value("Leave blank").run()
        self.assertFalse(app.exception)
        self.assertEqual(app.dataframe[-1].value["ChristmasAmnesty"].tolist(), ["", ""])
        self.assertEqual(app.dataframe[-1].value["NewEndo"].tolist(), ["", "Renewed"])

    def test_header_row_can_be_corrected_after_a_parse_error(self):
        app = self.start_app()
        app.file_uploader[0].set_value(("source.csv", SOURCE_CSV, "text/csv"))
        app.run()
        app.number_input[0].set_value(20).run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        app.number_input[0].set_value(1).run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertEqual(app.dataframe[-1].value.columns.tolist(), TEMPLATE_COLUMNS)


class AmnestyTests(unittest.TestCase):
    def source_and_output(self):
        source = pd.concat([read_source(SOURCE_CSV, "source.csv")] * 3, ignore_index=True)
        source["Area"] = ["MINDANAO", " mindanao ", "R3", "REGION 3", " R 3 ", "Other"]
        source["Christmas Amnesty"] = ["", "Eligible", " \t", "Yes", "0", "Eligible"]
        output = populate_template(source, suggest_mapping(source.columns))
        return source, output

    def test_blank_and_populated_tags_split_by_normalized_area(self):
        source, output = self.source_and_output()
        groups, excluded = split_amnesty(source, output, "Area", "Christmas Amnesty")
        self.assertEqual({label: len(frame) for label, frame in groups.items()}, {
            "AMNESTY NEGATIVE (MINDANAO)": 1,
            "AMNESTY (MINDANAO)": 1,
            "AMNESTY NEGATIVE (R3)": 1,
            "AMNESTY (R3)": 2,
        })
        self.assertEqual(excluded, 1)
        self.assertEqual(groups["AMNESTY (R3)"]["ChristmasAmnesty"].tolist(), ["Yes", "0"])
        self.assertEqual(groups["AMNESTY (R3)"]["chname"].tolist(), ["Jamie Example", "Alex Example"])
        self.assertTrue(all(frame.columns.tolist() == TEMPLATE_COLUMNS for frame in groups.values()))
        self.assertEqual(sum(len(frame) for frame in groups.values()) + excluded, len(source))

    def test_none_is_blank_and_zero_or_false_is_a_value(self):
        source = pd.DataFrame({"Area": ["R3"] * 4, "Christmas Amnesty": [None, "", 0, False]})
        output = pd.DataFrame({"marker": ["none", "empty", "zero", "false"]})
        groups, _ = split_amnesty(source, output, "Area", "Christmas Amnesty")
        self.assertEqual(groups["AMNESTY NEGATIVE (R3)"]["marker"].tolist(), ["none", "empty"])
        self.assertEqual(groups["AMNESTY (R3)"]["marker"].tolist(), ["zero", "false"])

    def test_source_status_controls_grouping_even_if_output_is_modified(self):
        source, output = self.source_and_output()
        source.index = [20, 30, 40, 50, 60, 70]
        output["ChristmasAmnesty"] = "Changed output"
        groups, _ = split_amnesty(source, output, "Area", "Christmas Amnesty")
        self.assertEqual(len(groups["AMNESTY NEGATIVE (MINDANAO)"]), 1)
        self.assertEqual(len(groups["AMNESTY NEGATIVE (R3)"]), 1)

    def test_missing_grouping_fields_are_not_assumed_to_be_negative(self):
        source, output = self.source_and_output()
        with self.assertRaises(ValidationError):
            split_amnesty(source, output, "missing", "Christmas Amnesty")
        with self.assertRaises(ValidationError):
            split_amnesty(source, output, "Area", "missing")

    def test_zip_has_four_plain_templates_with_correct_group_rows(self):
        source, output = self.source_and_output()
        groups, _ = split_amnesty(source, output, "Area", "Christmas Amnesty")
        exported_at = datetime(2026, 10, 13, 12, 30, 45, 123456, tzinfo=timezone.utc)
        downloads, archive = prepare_group_downloads(groups, "source.xlsx", exported_at)
        with ZipFile(io.BytesIO(archive)) as zip_file:
            self.assertEqual(len(zip_file.namelist()), 4)
            for label, download in downloads.items():
                self.assertTrue(download["filename"].startswith(label + "_source_20261013_203045_123456"))
                self.assertEqual(zip_file.read(download["filename"]), download["content"])
                sheet = load_workbook(io.BytesIO(download["content"])).active
                self.assertEqual([cell.value for cell in sheet[1]], TEMPLATE_COLUMNS)
                self.assertEqual(sheet.max_row - 1, len(groups[label]))
                self.assertTrue(all(not cell.has_style for row in sheet for cell in row))

    def test_known_source_headers_match_without_guessing_missing_identifiers(self):
        mapping = suggest_mapping(["email", "debtor_email", "name", "card_no", "OB (MASTERLIST}", "phone_1", "Phone 1", "account", "old_ic"])
        self.assertEqual(mapping["email"], "email")
        self.assertEqual(mapping["chname"], "name")
        self.assertEqual(mapping["cardnumber"], "card_no")
        self.assertEqual(mapping["OB"], "OB (MASTERLIST}")
        self.assertEqual(mapping["phone"], "Phone 1")
        self.assertIsNone(mapping["primary_key"])
        self.assertIsNone(mapping["indexoneuid"])

    def test_pending_enrichment_fields_stay_blank_and_phone_1_is_used(self):
        source = read_source(SOURCE_CSV, "source.csv").drop(columns=["Primary Key", "Index One UID", "Code", "DO Class", "Level", "Endo Date"])
        source = source.rename(columns={"Phone Number": "Phone 1"})
        source["phone_1"] = ["OTHER_NUMBER_1", "OTHER_NUMBER_2"]
        source["Area"] = ["MINDANAO", "REGION 3"]
        source["Christmas Amnesty"] = ["", "Eligible"]
        output = populate_template(source, suggest_mapping(source.columns))
        self.assertEqual(output["phone"].tolist(), source["Phone 1"].tolist())
        pending_fields = ["primary_key", "indexoneuid", "CODE", "doclass", "Level", "EndoDate"]
        self.assertTrue(output[pending_fields].eq("").all().all())
        app = AppTest.from_file(str(APP_PATH), default_timeout=20).run()
        app.file_uploader[0].set_value(("source.csv", source.to_csv(index=False).encode(), "text/csv"))
        app.run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertFalse(app.warning)
        self.assertEqual(len(app.get("download_button")), 3)
        downloads = app.session_state["email_group_downloads"]
        for label, download in downloads.items():
            sheet = load_workbook(io.BytesIO(download["content"])).active
            expected_phone = "09123456789" if "MINDANAO" in label else "09987654321"
            self.assertEqual(sheet["F2"].value, expected_phone)
            self.assertIsNone(sheet["B2"].value)
            self.assertIsNone(sheet["C2"].value)

    def test_ui_prepares_group_downloads_and_a_zip(self):
        source, _ = self.source_and_output()
        app = AppTest.from_file(str(APP_PATH), default_timeout=20).run()
        app.file_uploader[0].set_value(("source.csv", source.to_csv(index=False).encode(), "text/csv"))
        app.run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertEqual(len(app.get("download_button")), 5)
        self.assertEqual(len(app.session_state["email_group_downloads"]), 4)
        self.assertTrue(any("1 row(s) excluded" in warning.value for warning in app.warning))
        self.assertEqual(app.dataframe[-1].value.columns.tolist(), TEMPLATE_COLUMNS)
        next(button for button in app.button if button.label == "Clear source file").click().run()
        self.assertNotIn("email_group_downloads", app.session_state)
        self.assertFalse(app.get("download_button"))

    def test_ui_requires_the_source_amnesty_column_for_grouping(self):
        source, _ = self.source_and_output()
        source = source.drop(columns=["Christmas Amnesty"])
        app = AppTest.from_file(str(APP_PATH), default_timeout=20).run()
        app.file_uploader[0].set_value(("source.csv", source.to_csv(index=False).encode(), "text/csv"))
        app.run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        self.assertFalse(app.get("download_button"))
        self.assertTrue(any("Christmas Amnesty" in warning.value for warning in app.warning))


if __name__ == "__main__":
    unittest.main()
